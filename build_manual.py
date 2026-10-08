from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "PANDUAN-PENGGUNAAN.md"
OUTPUT = ROOT / "Panduan_Penggunaan_SIMAS_Gereja.docx"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', size='6'):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def set_run_font(run, name='Aptos', size=10.5, color='222222', bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold

def add_inline(paragraph, text):
    # Render the small amount of Markdown emphasis used in the guide.
    pattern = re.compile(r'(\*\*.*?\*\*|`.*?`)')
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith('**'):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name='Aptos Mono', size=9.5, color='174A8B')
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run)

def add_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for idx, value in enumerate(rows[0]):
        cell = table.rows[0].cells[idx]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(value.strip())
        set_run_font(run, size=9, color='FFFFFF', bold=True)
        set_cell_shading(cell, '174A8B')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        set_cell_margins(cell)
    for row_index, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value.strip())
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, 'F3F6FA')
            set_cell_border(cell)
            set_cell_margins(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Halaman ')
    set_run_font(run, size=8, color='667085')
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)

def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size, color in [('Heading 1', 16, '174A8B'), ('Heading 2', 12.5, '174A8B'), ('Heading 3', 11, '344054')]:
        style = styles[style_name]
        style.font.name = 'Aptos Display'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(13 if style_name == 'Heading 1' else 9)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    lines = SOURCE.read_text(encoding='utf-8').replace('\r\n', '\n').split('\n')
    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i].strip('\ufeff')
        if not line.strip():
            i += 1
            continue
        if line.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(7)
            for idx, row in enumerate(code):
                run = p.add_run(row + ('\n' if idx < len(code)-1 else ''))
                set_run_font(run, name='Aptos Mono', size=9, color='174A8B')
            i += 1
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(style='Title')
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[2:].strip())
            set_run_font(run, name='Aptos Display', size=24, color='111827', bold=True)
            first_heading = False
            i += 1
            continue
        heading = re.match(r'^(#{2,3})\s+(.*)$', line)
        if heading:
            level = 'Heading 1' if len(heading.group(1)) == 2 else 'Heading 2'
            p = doc.add_paragraph(style=level)
            add_inline(p, heading.group(2).strip())
            i += 1
            continue
        if line.startswith('|') and i + 1 < len(lines) and lines[i+1].strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                raw = lines[i].strip().strip('|')
                if not re.fullmatch(r'\s*:?[-]+:?\s*(\|\s*:?[-]+:?\s*)+\|?', raw):
                    table_lines.append([part.strip() for part in raw.split('|')])
                i += 1
            add_table(doc, table_lines)
            continue
        bullet = re.match(r'^[-*]\s+(.*)$', line)
        if bullet:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.18)
            add_inline(p, bullet.group(1))
            i += 1
            continue
        ordered = re.match(r'^(\d+)\.\s+(.*)$', line)
        if ordered:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Inches(0.18)
            add_inline(p, ordered.group(2))
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    # Keep metadata neutral and useful.
    doc.core_properties.title = 'Panduan Penggunaan SIMAS Gereja Aset Manajemen'
    doc.core_properties.subject = 'Panduan penggunaan aplikasi aset Paroki Pringwulung'
    doc.core_properties.author = 'SIMAS Gereja'
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == '__main__':
    build()
